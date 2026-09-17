#!/usr/bin/env python3
"""从月度汇总表生成 V2 仪表板数据。"""

from __future__ import annotations

import argparse
import json
import math
from datetime import datetime
from pathlib import Path

import openpyxl


METRIC_GROUPS = [
    {
        "id": "basic",
        "title": "基础使用",
        "description": "企微开通、登录与智慧导购使用规模。",
        "rows": [4, 5, 6],
    },
    {
        "id": "active",
        "title": "活跃度",
        "description": "全体与三类人员的日活、月活趋势。",
        "rows": [7, 8, 9, 10],
    },
    {
        "id": "features",
        "title": "功能使用",
        "description": "厨房健康检测与选购指南的核心使用指标。",
        "rows": [11, 12, 13, 14, 15],
    },
    {
        "id": "moments",
        "title": "触达运营",
        "description": "朋友圈创建、执行与触达效果。",
        "rows": [16, 17, 18, 19, 20],
    },
    {
        "id": "welcome",
        "title": "欢迎语累计",
        "description": "欢迎语触达、打开与线索累计效果。",
        "rows": [27, 28, 29, 30, 31, 32, 33],
    },
]

HIGHLIGHT_ROWS = [4, 5, 6, 10]
PERCENT_ROWS = {19, 30, 32}


def clean_value(value):
    if value is None:
        return None
    if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
        return None
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def col_to_year_month(sheet, col):
    month_label = sheet.cell(3, col).value
    if not month_label or not str(month_label).endswith("月"):
        return None

    year_label = sheet.cell(2, col).value
    if year_label is None:
        # Excel 合并单元格下，年份只在第一个月份列出现。
        for prev_col in range(col - 1, 0, -1):
            prev_year = sheet.cell(2, prev_col).value
            if prev_year is not None:
                year_label = prev_year
                break

    year_text = str(year_label).replace("年", "")
    month = int(str(month_label).replace("月", ""))
    year = 2000 + int(year_text)
    return {
        "key": f"{year}-{month:02d}",
        "label": f"{str(year_label)}{month}月",
        "shortLabel": f"{month}月",
        "year": year,
        "month": month,
    }


def metric_from_row(sheet, row, months, latest_col):
    name = sheet.cell(row, 1).value
    values = [clean_value(sheet.cell(row, month["col"]).value) for month in months]
    latest = clean_value(sheet.cell(row, latest_col).value)
    mom = clean_value(sheet.cell(row, latest_col + 2).value)
    yoy = clean_value(sheet.cell(row, latest_col + 1).value)
    return {
        "id": f"row-{row}",
        "row": row,
        "name": name,
        "unit": "%" if row in PERCENT_ROWS else "",
        "format": "percent" if row in PERCENT_ROWS else "number",
        "values": values,
        "latest": latest,
        "mom": mom,
        "yoy": yoy,
    }


def build_dashboard_data(workbook_path: Path):
    wb = openpyxl.load_workbook(workbook_path, data_only=True)
    sheet = wb["Sheet1"]

    months = []
    for col in range(2, sheet.max_column + 1):
        item = col_to_year_month(sheet, col)
        if item:
            item["col"] = col
            months.append(item)

    if not months:
        raise ValueError("未在 Sheet1 第 3 行找到月份列")

    latest_col = months[-1]["col"]
    latest_month = months[-1]

    all_metrics = {
        row: metric_from_row(sheet, row, months, latest_col)
        for group in METRIC_GROUPS
        for row in group["rows"]
    }

    groups = []
    for group in METRIC_GROUPS:
        groups.append(
            {
                "id": group["id"],
                "title": group["title"],
                "description": group["description"],
                "metrics": [all_metrics[row] for row in group["rows"]],
            }
        )

    highlights = [all_metrics[row] for row in HIGHLIGHT_ROWS]

    wb.close()

    return {
        "meta": {
            "title": "企微智慧导购月度数据后台 V2",
            "source": workbook_path.name,
            "generatedAt": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "latestMonth": {
                key: latest_month[key]
                for key in ["key", "label", "shortLabel", "year", "month"]
            },
            "notes": [
                "数据来自月度汇总数据(终）.xlsx。",
                "日活口径：按天去重业务员编码后求平均，并取整数。",
                "月活口径：当月去重业务员编码。",
                "岗位口径边界：2026年3—7月保留历史已发布结果，不追溯调整。",
                "2026年8月起三类人员岗位：厨电顾问、厨电顾问组长、家装专厅厨电顾问、店长、客户经理、家装客户经理-门店、客户经理组长。",
                "跨越2026年8月口径切换边界时，三类人员日活、月活的环比和同比需结合岗位范围变化解读。",
                "当前版本仅展示汇总表中稳定维护的指标；未持续维护的旧模块暂不展示。",
            ],
        },
        "months": [
            {key: month[key] for key in ["key", "label", "shortLabel", "year", "month"]}
            for month in months
        ],
        "highlights": highlights,
        "groups": groups,
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path, help="月度汇总 Excel 文件路径")
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("v2/data/dashboard-data.json"),
        help="输出 JSON 路径",
    )
    args = parser.parse_args()

    data = build_dashboard_data(args.workbook)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        json.dumps(data, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(f"已生成 {args.output}")
    print(f"最新月份：{data['meta']['latestMonth']['label']}")


if __name__ == "__main__":
    main()
